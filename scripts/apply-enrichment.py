"""Parse enrich-descriptions and enrich-radars JSON responses, update Excel with rich content.

Approach:
1. Load Sites + Radars from Excel
2. Load Sources to build notebook_uuid -> SRC-NNNN mapping
3. Walk through site-description batch JSONs:
   - Parse answer text for SITE_BEGIN / SITE_END blocks
   - Remap [1] [2] citations to [SRC-NNNN] using references array
   - Update Sites sheet description field + a new citations field
4. Walk through radar-description batch JSONs similarly
5. Save Excel
"""
import os
import re
import json
import glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL = os.path.join(BASE, "data", "us_missile_range_data.xlsx")
DESC_DIRS = [
    os.path.join(BASE, "data", "notebook-responses", "site-descriptions"),
    os.path.join(BASE, "data", "notebook-responses", "site-descriptions-v2"),
]
RADAR_DIRS = [
    os.path.join(BASE, "data", "notebook-responses", "radar-descriptions"),
    os.path.join(BASE, "data", "notebook-responses", "radar-descriptions-v2"),
]


def load_uuid_to_src(wb):
    sheet = wb["Sources"]
    headers = [c.value for c in sheet[1]]
    uuid_col = headers.index("notebook_uuid") + 1
    sid_col = headers.index("source_id") + 1
    out = {}
    for row in range(2, sheet.max_row + 1):
        uuid = sheet.cell(row=row, column=uuid_col).value
        sid = sheet.cell(row=row, column=sid_col).value
        if uuid and sid:
            out[uuid] = sid
    return out


def remap_citations(text: str, references: list, uuid_to_src: dict) -> tuple:
    """Replace [1] [2] in text with [SRC-NNNN]. Return (new_text, list_of_src_ids)."""
    # Build citation_number -> source_id mapping (deduped, ordered)
    num_to_src = {}
    for ref in references:
        n = ref.get("citation_number")
        nb_uuid = ref.get("source_id")
        if n and nb_uuid:
            src_id = uuid_to_src.get(nb_uuid)
            if src_id:
                num_to_src[n] = src_id

    used_src_ids = []

    def replace_match(m):
        # m.group(1) is what's inside the brackets, like "1", "2-4", "1, 3, 5"
        inside = m.group(1)
        parts = re.split(r"[,\s]+", inside)
        replaced = []
        for p in parts:
            p = p.strip()
            if not p:
                continue
            if "-" in p:
                # Range like 2-4
                try:
                    a, b = p.split("-", 1)
                    for n in range(int(a), int(b) + 1):
                        if n in num_to_src:
                            src = num_to_src[n]
                            if src not in used_src_ids:
                                used_src_ids.append(src)
                            replaced.append(src)
                except ValueError:
                    replaced.append(p)
            else:
                try:
                    n = int(p)
                    if n in num_to_src:
                        src = num_to_src[n]
                        if src not in used_src_ids:
                            used_src_ids.append(src)
                        replaced.append(src)
                    else:
                        replaced.append(p)
                except ValueError:
                    replaced.append(p)
        return "[" + ", ".join(replaced) + "]" if replaced else m.group(0)

    new_text = re.sub(r"\[([\d\s,\-]+)\]", replace_match, text)
    return new_text, used_src_ids


def parse_site_descriptions(answer_text: str) -> dict:
    """Parse 'SITE_BEGIN: name\n<body>\nSITE_END' blocks. Returns {name -> body}."""
    out = {}
    # Split on SITE_BEGIN
    pattern = re.compile(r"SITE_BEGIN[:\s]*(.+?)\n(.*?)SITE_END", re.DOTALL | re.IGNORECASE)
    for m in pattern.finditer(answer_text):
        name = m.group(1).strip().rstrip(":").strip()
        body = m.group(2).strip()
        if name and body and "no detailed information" not in body.lower():
            out[name] = body
    return out


def parse_radar_descriptions(answer_text: str) -> dict:
    """Parse 'RADAR_BEGIN: name | site\n<body>\nRADAR_END' blocks. Returns {(name, site) -> body}."""
    out = {}
    pattern = re.compile(r"RADAR_BEGIN[:\s]*(.+?)\n(.*?)RADAR_END", re.DOTALL | re.IGNORECASE)
    for m in pattern.finditer(answer_text):
        header = m.group(1).strip()
        body = m.group(2).strip()
        if "|" in header:
            name, site = header.split("|", 1)
            name = name.strip()
            site = site.strip()
        else:
            name = header
            site = ""
        if name and body and "no detailed information" not in body.lower():
            out[(name.lower(), site.lower())] = body
    return out


def ensure_column(sheet, header_name, width=40):
    headers = [c.value for c in sheet[1]]
    if header_name in headers:
        return headers.index(header_name) + 1
    new_col = sheet.max_column + 1
    sheet.cell(row=1, column=new_col, value=header_name)
    sheet.column_dimensions[openpyxl.utils.get_column_letter(new_col)].width = width
    return new_col


def update_site_descriptions(wb, uuid_to_src):
    sheet = wb["Sites"]
    headers = [c.value for c in sheet[1]]
    name_col = headers.index("site_name") + 1
    desc_col = headers.index("description") + 1
    cite_col = ensure_column(sheet, "citations", width=60)

    # Build name -> row map
    name_to_row = {}
    for row in range(2, sheet.max_row + 1):
        name = sheet.cell(row=row, column=name_col).value
        if name:
            name_to_row[name.lower().strip()] = row

    updated = 0
    files = []
    for d in DESC_DIRS:
        files.extend(sorted(glob.glob(os.path.join(d, "*.json"))))
    print(f"  Processing {len(files)} site description batches...")
    for fname in files:
        try:
            with open(fname, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception as e:
            print(f"    Skip {os.path.basename(fname)}: {e}")
            continue
        answer = data.get("answer", "")
        references = data.get("references", [])
        if not answer:
            continue

        site_blocks = parse_site_descriptions(answer)
        for site_name, body in site_blocks.items():
            new_text, src_ids = remap_citations(body, references, uuid_to_src)
            # Find row
            key = site_name.lower().strip()
            row = name_to_row.get(key)
            if not row:
                # Try fuzzy: contains
                for k, r in name_to_row.items():
                    if k in key or key in k:
                        if min(len(k), len(key)) >= 5:
                            row = r
                            break
            if row:
                sheet.cell(row=row, column=desc_col, value=new_text)
                sheet.cell(row=row, column=cite_col, value=", ".join(src_ids))
                updated += 1
    print(f"  Updated {updated} site descriptions")


def update_radar_descriptions(wb, uuid_to_src):
    sites_sheet = wb["Sites"]
    radars_sheet = wb["Radars"]

    s_headers = [c.value for c in sites_sheet[1]]
    s_id_col = s_headers.index("site_id") + 1
    s_name_col = s_headers.index("site_name") + 1
    id_to_sname = {}
    sname_to_id = {}
    for row in range(2, sites_sheet.max_row + 1):
        sid = sites_sheet.cell(row=row, column=s_id_col).value
        name = sites_sheet.cell(row=row, column=s_name_col).value
        if sid:
            norm = (name or "").lower().strip()
            id_to_sname[sid] = norm
            sname_to_id[norm] = sid

    r_headers = [c.value for c in radars_sheet[1]]
    r_id_col = r_headers.index("radar_id") + 1
    r_name_col = r_headers.index("radar_name") + 1
    r_site_col = r_headers.index("site_id") + 1
    r_desc_col = r_headers.index("public_description") + 1
    r_cite_col = ensure_column(radars_sheet, "citations", width=60)

    # Find highest existing RAD-NNNN
    max_rad = 0
    for row in range(2, radars_sheet.max_row + 1):
        rid = radars_sheet.cell(row=row, column=r_id_col).value
        if rid and rid.startswith("RAD-"):
            try:
                n = int(rid[4:])
                if n > max_rad:
                    max_rad = n
            except ValueError:
                pass

    # (name_lower, site_name_lower) -> row
    radar_index = {}
    for row in range(2, radars_sheet.max_row + 1):
        rname = radars_sheet.cell(row=row, column=r_name_col).value or ""
        sid = radars_sheet.cell(row=row, column=r_site_col).value
        site_name = id_to_sname.get(sid, "")
        radar_index[(rname.lower().strip(), site_name)] = row

    updated = 0
    inserted = 0
    files = []
    for d in RADAR_DIRS:
        files.extend(sorted(glob.glob(os.path.join(d, "*.json"))))
    print(f"  Processing {len(files)} radar description batches...")
    for fname in files:
        try:
            with open(fname, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception:
            continue
        answer = data.get("answer", "")
        references = data.get("references", [])
        if not answer:
            continue
        blocks = parse_radar_descriptions(answer)
        for (radar_name, site_name), body in blocks.items():
            new_text, src_ids = remap_citations(body, references, uuid_to_src)
            # Match radar row
            row = radar_index.get((radar_name, site_name))
            if not row:
                # Fuzzy: same site_name, radar name contains
                for (rn, sn), r in radar_index.items():
                    if sn == site_name and (radar_name in rn or rn in radar_name) and min(len(rn), len(radar_name)) >= 4:
                        row = r
                        break
            if row:
                radars_sheet.cell(row=row, column=r_desc_col, value=new_text)
                radars_sheet.cell(row=row, column=r_cite_col, value=", ".join(src_ids))
                updated += 1
            else:
                # Need to insert a new radar row, but only if we can map site_name -> site_id
                sid = sname_to_id.get(site_name)
                if not sid:
                    # Try fuzzy site match
                    for sn, s_id in sname_to_id.items():
                        if site_name and (site_name in sn or sn in site_name) and min(len(sn), len(site_name)) >= 5:
                            sid = s_id
                            break
                if not sid:
                    continue
                max_rad += 1
                new_row = [None] * radars_sheet.max_column
                new_row[r_id_col - 1] = f"RAD-{max_rad:04d}"
                new_row[r_site_col - 1] = sid
                new_row[r_name_col - 1] = radar_name.title() if radar_name else ""
                new_row[r_desc_col - 1] = new_text
                # operational_status default to Unknown
                if "operational_status" in r_headers:
                    new_row[r_headers.index("operational_status")] = "Unknown"
                if "radar_type" in r_headers:
                    new_row[r_headers.index("radar_type")] = "Unknown"
                if "purpose" in r_headers:
                    new_row[r_headers.index("purpose")] = "See description"
                if "confidence_level" in r_headers:
                    new_row[r_headers.index("confidence_level")] = "Medium"
                if "record_status" in r_headers:
                    new_row[r_headers.index("record_status")] = "Published"
                if "source_id" in r_headers and src_ids:
                    new_row[r_headers.index("source_id")] = src_ids[0]
                new_row[r_cite_col - 1] = ", ".join(src_ids)
                radars_sheet.append(new_row)
                # Index the newly added radar
                radar_index[(radar_name.lower().strip(), site_name)] = radars_sheet.max_row
                inserted += 1
    print(f"  Updated {updated} radar descriptions, inserted {inserted} new radars")


def main():
    print(f"Loading {EXCEL}")
    wb = openpyxl.load_workbook(EXCEL)
    uuid_to_src = load_uuid_to_src(wb)
    print(f"  {len(uuid_to_src)} source UUIDs mapped to SRC-NNNN")

    update_site_descriptions(wb, uuid_to_src)
    update_radar_descriptions(wb, uuid_to_src)

    wb.save(EXCEL)
    print(f"Saved {EXCEL}")


if __name__ == "__main__":
    main()
