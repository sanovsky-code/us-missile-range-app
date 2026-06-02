"""Debug script: show which radar site_names from notebook don't match Sites sheet."""
import glob
import re
import difflib
import openpyxl
from collections import defaultdict

wb = openpyxl.load_workbook('data/us_missile_range_data.xlsx')
sites = wb['Sites']
headers = [c.value for c in sites[1]]
nc = headers.index('site_name') + 1
cc = headers.index('country') + 1

country_sites = defaultdict(set)
for row in range(2, sites.max_row + 1):
    name = sites.cell(row=row, column=nc).value
    country = sites.cell(row=row, column=cc).value
    if name and country:
        country_sites[country].add(name)

unmatched = defaultdict(list)
for fname in glob.glob('data/notebook-responses/radars-*.txt'):
    with open(fname, 'r', encoding='utf-8') as f:
        raw = f.read()
    chunks = []
    current = None
    for line in raw.split('\n'):
        if line.startswith('RADAR|'):
            if current:
                chunks.append(current)
            current = line
        elif current and not line.startswith(('Answer:', 'New conversation:')):
            current += ' ' + line.strip()
    if current:
        chunks.append(current)
    for line in chunks:
        line = re.sub(r'\s+', ' ', line).strip()
        parts = line.split('|')
        if len(parts) < 13:
            continue
        country = parts[1].strip()
        if country in ('United States of America', 'United States'):
            country = 'USA'
        if country == 'United Kingdom':
            country = 'UK'
        site_name = parts[2].strip()
        if not site_name:
            continue
        existing = country_sites.get(country, set())
        if site_name in existing:
            continue
        match = difflib.get_close_matches(site_name, list(existing), n=1, cutoff=0.75)
        if match:
            continue
        unmatched[country].append(site_name)

for country in sorted(unmatched.keys()):
    print(f'\n=== {country}: {len(unmatched[country])} unmatched ===')
    for s in sorted(set(unmatched[country])):
        print(f'  - {s}')
