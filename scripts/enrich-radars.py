"""Query NotebookLM in batches for detailed radar descriptions with inline citations.

For each site that has radars, query the notebook for rich details on each radar.
"""
import os
import json
import subprocess
import openpyxl
from concurrent.futures import ThreadPoolExecutor, as_completed
from collections import defaultdict

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL = os.path.join(BASE, "data", "us_missile_range_data.xlsx")
RESPONSES_DIR = os.path.join(BASE, "data", "notebook-responses")
RADAR_DIR = os.path.join(RESPONSES_DIR, "radar-descriptions")
os.makedirs(RADAR_DIR, exist_ok=True)

BATCH_SIZE = 6
MAX_PARALLEL = 5


def load_radars_with_sites():
    wb = openpyxl.load_workbook(EXCEL)
    sites_sheet = wb["Sites"]
    radars_sheet = wb["Radars"]

    s_headers = [c.value for c in sites_sheet[1]]
    site_id_col = s_headers.index("site_id") + 1
    name_col = s_headers.index("site_name") + 1
    country_col = s_headers.index("country") + 1

    site_info = {}
    for row in range(2, sites_sheet.max_row + 1):
        sid = sites_sheet.cell(row=row, column=site_id_col).value
        name = sites_sheet.cell(row=row, column=name_col).value
        country = sites_sheet.cell(row=row, column=country_col).value
        if sid:
            site_info[sid] = {"name": name, "country": country or ""}

    r_headers = [c.value for c in radars_sheet[1]]
    radar_id_col = r_headers.index("radar_id") + 1
    radar_site_col = r_headers.index("site_id") + 1
    radar_name_col = r_headers.index("radar_name") + 1
    radar_model_col = r_headers.index("radar_model") + 1

    radars = []
    for row in range(2, radars_sheet.max_row + 1):
        rid = radars_sheet.cell(row=row, column=radar_id_col).value
        sid = radars_sheet.cell(row=row, column=radar_site_col).value
        rname = radars_sheet.cell(row=row, column=radar_name_col).value
        rmodel = radars_sheet.cell(row=row, column=radar_model_col).value
        if rid and sid:
            radars.append({
                "radar_id": rid, "site_id": sid,
                "name": rname, "model": rmodel,
                "site_name": site_info.get(sid, {}).get("name", ""),
                "country": site_info.get(sid, {}).get("country", ""),
            })
    return radars


def build_batches(radars):
    by_country = defaultdict(list)
    for r in radars:
        by_country[r["country"]].append(r)
    batches = []
    for country, country_radars in by_country.items():
        for i in range(0, len(country_radars), BATCH_SIZE):
            batches.append((country, country_radars[i:i + BATCH_SIZE]))
    return batches


def make_prompt(country: str, radars: list) -> str:
    list_text = "\n".join(
        f"- {r['name']}" + (f" ({r['model']})" if r['model'] else "") + f" at {r['site_name']}"
        for r in radars
    )
    return f"""For each of the following radar systems in {country}, write a detailed description (150-250 words each) based on the sources. Cite inline using [1], [2] style.

Radars:
{list_text}

Format EXACTLY like this (separated by ###):

RADAR_BEGIN: <radar name> | <site name>
<detailed description with [1] [2] inline citations covering: type, frequency band, manufacturer, year built, capabilities, purpose, dimensions/power if known, current status, notable operations/upgrades>
RADAR_END

###

Begin immediately with RADAR_BEGIN:"""


def run_query(batch_idx: int, country: str, radars: list) -> dict:
    out_file = os.path.join(RADAR_DIR, f"batch-{batch_idx:03d}-{country.replace(' ', '_')[:20]}.json")
    if os.path.exists(out_file):
        try:
            with open(out_file, "r", encoding="utf-8") as f:
                data = json.load(f)
            if data.get("answer"):
                return {"batch_idx": batch_idx, "country": country, "status": "cached"}
        except Exception:
            pass
    prompt = make_prompt(country, radars)
    try:
        proc = subprocess.run(
            ["python", "-m", "notebooklm", "ask", "--new", "-y", "--json", prompt],
            capture_output=True, text=True, timeout=600,
            encoding="utf-8", errors="replace",
        )
        with open(out_file, "w", encoding="utf-8") as f:
            f.write(proc.stdout)
        return {"batch_idx": batch_idx, "country": country,
                "status": "ok" if proc.returncode == 0 else "error"}
    except subprocess.TimeoutExpired:
        return {"batch_idx": batch_idx, "country": country, "status": "timeout"}


def main():
    radars = load_radars_with_sites()
    print(f"Loaded {len(radars)} radars")
    batches = build_batches(radars)
    print(f"Built {len(batches)} batches (max {BATCH_SIZE} radars each)")

    with ThreadPoolExecutor(max_workers=MAX_PARALLEL) as ex:
        futures = {ex.submit(run_query, i, country, rs): (i, country)
                   for i, (country, rs) in enumerate(batches)}
        for fut in as_completed(futures):
            res = fut.result()
            print(f"  [{res['status']:8}] batch {res['batch_idx']:3} ({res['country']})")
    print(f"\nDone. Saved to {RADAR_DIR}")


if __name__ == "__main__":
    main()
